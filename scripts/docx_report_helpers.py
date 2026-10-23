"""
Script to generate the complete, production-grade College Project Implementation Report (MCA-4004)
for IIIT Bhopal adhering strictly to the official MCA 2026 guidelines.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove existing shading if present
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Set inner cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="cbd5e1", sz="4", val="single"):
    """Apply elegant, subtle borders to table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_table_headers(table, col_names, col_widths=None):
    """Format table header row with deep navy theme and white bold text."""
    hdr_cells = table.rows[0].cells
    for i, name in enumerate(col_names):
        hdr_cells[i].text = name
        set_cell_shading(hdr_cells[i], "1e3a8a")
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=160, right=160)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    
    # Repeat header row across pages
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def populate_table_rows(table, data, col_widths=None, align_left_col0=True):
    """Fill table data with alternating row colors."""
    for r_idx, row_data in enumerate(data):
        if len(table.rows) <= r_idx + 1:
            row = table.add_row()
        else:
            row = table.rows[r_idx + 1]
        bg_color = "f8fafc" if r_idx % 2 == 1 else "ffffff"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = str(val)
            set_cell_shading(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            if c_idx == 0 and align_left_col0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
                
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

def add_styled_heading(doc, text, level):
    """Add heading strictly conforming to college guidelines:
    Level 1: 16pt Bold Times New Roman (Heading 1)
    Level 2: 14pt Bold Times New Roman (Heading 2)
    Level 3: 12pt Bold Times New Roman (Heading 3)
    """
    style_name = f"Heading {level}"
    try:
        p = doc.add_paragraph(style=style_name)
    except Exception:
        p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.bold = True
    
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(15, 23, 42) # Dark Navy
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(30, 41, 59) # Slate Navy
    elif level == 3:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(51, 65, 85) # Slate
    return p

def setup_header_and_footer(doc):
    """Add formal running header and page numbering footer, unlinking title page."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for sec in doc.sections:
        sec.different_first_page_header_footer = True
        
        # Cover page clean
        if sec.first_page_header.paragraphs:
            sec.first_page_header.paragraphs[0].text = ""
        if sec.first_page_footer.paragraphs:
            sec.first_page_footer.paragraphs[0].text = ""
            
        # Running Header
        hdr_p = sec.header.paragraphs[0]
        hdr_p.text = ""
        hdr_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hdr_run = hdr_p.add_run("Department of Information Technology, IIIT Bhopal | MCA-4004 Project Report")
        hdr_run.font.name = "Times New Roman"
        hdr_run.font.size = Pt(8.5)
        hdr_run.italic = True
        hdr_run.font.color.rgb = RGBColor(100, 116, 139)
        
        # Running Footer (Centered Page X)
        ftr_p = sec.footer.paragraphs[0]
        ftr_p.text = ""
        ftr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ftr_run_pre = ftr_p.add_run("Page ")
        ftr_run_pre.font.name = "Times New Roman"
        ftr_run_pre.font.size = Pt(9.5)
        ftr_run_pre.font.color.rgb = RGBColor(100, 116, 139)
        
        ftr_run_pg = ftr_p.add_run()
        ftr_run_pg.font.name = "Times New Roman"
        ftr_run_pg.font.size = Pt(9.5)
        ftr_run_pg.font.color.rgb = RGBColor(100, 116, 139)
        
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = "PAGE"
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'separate')
        fldChar3 = OxmlElement('w:fldChar')
        fldChar3.set(qn('w:fldCharType'), 'end')
        
        r_node = ftr_run_pg._r
        r_node.append(fldChar1)
        r_node.append(instrText)
        r_node.append(fldChar2)
        r_node.append(fldChar3)


def add_body_p(doc, text, bold_prefix="", italic=False):
    """Add standard 12pt Times New Roman body text with 1.15 line spacing."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
        
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.italic = italic
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet_p(doc, text, bold_prefix=""):
    """Add bullet paragraph."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_figure_with_caption(doc, img_path, caption_text, width=Inches(5.8)):
    """Add figure centered with 11pt Times New Roman caption."""
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        p_cap.paragraph_format.keep_with_next = False
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(11)
        r_cap.italic = True
        r_cap.font.color.rgb = RGBColor(71, 85, 105)
    else:
        print(f"[WARN] Figure image not found: {img_path}")

def add_table_caption(doc, caption_text):
    """Add formal table caption before table."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(caption_text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = RGBColor(30, 41, 59)

def add_code_block(doc, code_text):
    """Add shaded monospaced box for algorithms or code snippets."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    set_cell_shading(c, "f1f5f9") # Light Slate Gray
    set_cell_margins(c, top=140, bottom=140, left=180, right=180)
    set_table_borders(tbl, color="cbd5e1", sz="6", val="single")
    
    lines = code_text.strip().split("\n")
    for i, line in enumerate(lines):
        clean_line = "".join(ch for ch in line if ch == "\t" or ord(ch) >= 32)
        if i == 0:
            p = c.paragraphs[0]
        else:
            p = c.add_paragraph()
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(clean_line)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(15, 23, 42)
    
    # spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)

print("Helper definitions loaded successfully.")
